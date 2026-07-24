"""
=========================================================================
 XAMPP & IIS Directory Scanner
=========================================================================
Aplikasi desktop (Windows) untuk melakukan audit direktori server lokal
(XAMPP htdocs & IIS wwwroot), mengidentifikasi project/aplikasi web di
dalamnya (Laravel, WordPress, Node.js, ASP.NET, PHP native, dll), lalu
meng-export hasilnya ke file Excel (.xlsx) yang sudah diformat rapi.

Cara Install Dependensi:
    pip install customtkinter pandas openpyxl

Cara Menjalankan:
    python xampp_iis_scanner.py

Catatan:
- Aplikasi ini didesain untuk berjalan di Windows (path seperti
  C:\\xampp\\htdocs, C:\\inetpub\\wwwroot). Di OS lain, path default
  tidak akan ditemukan, tetapi fitur "Tambah Path Manual" tetap bisa
  dipakai untuk memindai folder apa saja.
=========================================================================
"""

import os
import time
import queue
import threading
from datetime import datetime

import customtkinter as ctk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# =========================================================================
# KONFIGURASI & KONSTANTA
# =========================================================================

APP_TITLE = "XAMPP & IIS Directory Scanner by Nova"
APP_SUBTITLE = "Audit direktori server lokal & export hasil ke Excel"
WINDOW_SIZE = "1320x860"

# Folder berat yang di-skip agar proses scan ukuran folder tetap cepat.
# Folder ini tetap dihitung sebagai 1 subfolder, tapi isinya tidak
# ditelusuri lebih dalam (tidak dihitung ukuran & jumlah file-nya).
SKIP_DIR_NAMES = {"node_modules", "vendor", ".git"}

# Daftar path default yang otomatis diperiksa saat aplikasi dibuka.
# Format: (label tampilan, path)
DEFAULT_PATHS = [
    ("XAMPP Default (C:\\xampp\\htdocs)", r"C:\xampp\htdocs"),
    ("XAMPP Drive D (D:\\xampp\\htdocs)", r"D:\xampp\htdocs"),
    ("XAMPP8 Drive D (D:\\xampp8\\htdocs)", r"D:\xampp8\htdocs"),
    ("XAMPP8 Drive C (C:\\xampp8\\htdocs)", r"C:\xampp8\htdocs"),
    ("IIS Default (C:\\inetpub\\wwwroot)", r"C:\inetpub\wwwroot"),
]

EXPORT_COLUMNS = [
    "No",
    "Nama Project",
    "Sumber Server",
    "Full Path",
    "Framework/Tech",
    "Ukuran (MB)",
    "Terakhir Dimodifikasi",
    "Jumlah File",
]

PREVIEW_COLUMNS = [
    "No",
    "Nama Project",
    "Sumber Server",
    "Framework/Tech",
    "Ukuran",
    "Terakhir Dimodifikasi",
    "Jumlah File",
    "Full Path",
]


# =========================================================================
# LOGIKA SCANNER (murni logic, tidak menyentuh GUI, supaya mudah diuji)
# =========================================================================

def derive_source_label(base_path: str) -> str:
    """Menentukan label sumber server berdasarkan isi path (XAMPP/IIS/Manual)."""
    p = base_path.lower()
    if "xampp" in p:
        tag = "XAMPP"
    elif "inetpub" in p or "iis" in p:
        tag = "IIS"
    else:
        tag = "Manual"
    return f"{tag} ({base_path})"


def detect_tech_stack(project_path: str) -> str:
    """Mendeteksi framework/tech stack sebuah folder project berdasarkan
    keberadaan file penanda (marker file) di root folder tersebut."""

    def has(name: str) -> bool:
        return os.path.isfile(os.path.join(project_path, name))

    # 1. WordPress
    if has("wp-config.php"):
        return "WordPress"

    # 2. Laravel (artisan pasti Laravel)
    if has("artisan"):
        return "Laravel"

    # 3. composer.json -> cek isinya untuk memastikan Laravel atau PHP biasa
    if has("composer.json"):
        try:
            with open(os.path.join(project_path, "composer.json"),
                       "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            if "laravel/framework" in content:
                return "Laravel"
        except OSError:
            pass
        return "PHP (Composer)"

    # 4. package.json -> Node.js / React / Vue / Next / Angular
    if has("package.json"):
        try:
            with open(os.path.join(project_path, "package.json"),
                       "r", encoding="utf-8", errors="ignore") as f:
                content = f.read().lower()
            if '"next"' in content:
                return "Node.js (Next.js)"
            if '"react"' in content:
                return "Node.js (React)"
            if '"vue"' in content:
                return "Node.js (Vue)"
            if '"@angular/core"' in content:
                return "Node.js (Angular)"
        except OSError:
            pass
        return "Node.js"

    # 5. web.config -> ASP.NET / IIS Configured
    if has("web.config"):
        return "ASP.NET / IIS Configured"

    # 6. index.php -> Native PHP
    if has("index.php"):
        return "Native PHP"

    # 7. index.html -> Native HTML
    if has("index.html"):
        return "Native HTML"

    # 8. Tidak ada file utama
    return "Folder / Asset Biasa"


def scan_folder_stats(path: str, skip_dirs: set):
    """Menghitung total ukuran (bytes), jumlah file, jumlah subfolder, dan
    waktu modifikasi terakhir dari sebuah folder, sambil melewati (skip)
    folder berat seperti node_modules & vendor agar proses cepat."""

    total_size = 0
    file_count = 0
    folder_count = 0
    try:
        max_mtime = os.path.getmtime(path)
    except OSError:
        max_mtime = time.time()

    for root, dirs, files in os.walk(path):
        # Pisahkan subfolder yang harus di-skip dari yang ditelusuri lebih dalam
        skipped = [d for d in dirs if d.lower() in skip_dirs]
        dirs[:] = [d for d in dirs if d.lower() not in skip_dirs]

        # Folder yang di-skip tetap dihitung sebagai subfolder, tapi isinya tidak dihitung
        folder_count += len(dirs) + len(skipped)

        for fname in files:
            file_count += 1
            fpath = os.path.join(root, fname)
            try:
                st = os.stat(fpath)
                total_size += st.st_size
                if st.st_mtime > max_mtime:
                    max_mtime = st.st_mtime
            except OSError:
                continue

    return total_size, file_count, folder_count, max_mtime


def format_size_human(size_mb: float) -> str:
    """Format ukuran MB menjadi string yang enak dibaca (MB atau GB)."""
    if size_mb >= 1024:
        return f"{size_mb / 1024:.2f} GB"
    return f"{size_mb:.2f} MB"


def export_results_to_excel(rows: list, filepath: str):
    """Export list of dict hasil scan ke file .xlsx dengan header berwarna
    biru gelap + teks putih tebal, dan lebar kolom auto-fit."""

    # Susun DataFrame sesuai urutan kolom export yang diminta
    export_rows = []
    for r in rows:
        export_rows.append({
            "No": r["No"],
            "Nama Project": r["Nama Project"],
            "Sumber Server": r["Sumber Server"],
            "Full Path": r["Full Path"],
            "Framework/Tech": r["Framework/Tech"],
            "Ukuran (MB)": r["Ukuran (MB)"],
            "Terakhir Dimodifikasi": r["Terakhir Dimodifikasi"],
            "Jumlah File": r["Jumlah File"],
        })

    df = pd.DataFrame(export_rows, columns=EXPORT_COLUMNS)

    with pd.ExcelWriter(filepath, engine="openpyxl") as writer:
        sheet_name = "Hasil Scan"
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        ws = writer.sheets[sheet_name]

        # Styling header: biru gelap, teks putih tebal, rata tengah
        header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_align

        # Auto-fit lebar kolom berdasarkan panjang isi terpanjang
        for col_idx, col_name in enumerate(EXPORT_COLUMNS, start=1):
            max_len = len(str(col_name))
            for row in export_rows:
                val_len = len(str(row.get(col_name, "")))
                if val_len > max_len:
                    max_len = val_len
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = min(max_len + 4, 80)

        ws.freeze_panes = "A2"
        ws.row_dimensions[1].height = 22


# =========================================================================
# APLIKASI GUI (CustomTkinter)
# =========================================================================

class ScannerApp(ctk.CTk):
    def __init__(self):
        super().__init__()

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        self.title(APP_TITLE)
        self.geometry(WINDOW_SIZE)
        self.minsize(1100, 700)

        # State internal
        self.path_rows = {}          # row_id -> {"var", "path", "label", "removable", "frame"}
        self._path_row_counter = 0
        self.scan_results = []       # list of dict hasil scan (setelah selesai)
        self.log_queue = queue.Queue()
        self.is_scanning = False

        self._build_header()
        self._build_control_panel()
        self._build_progress_zone()
        self._build_result_zone()
        self._build_action_zone()

        self._populate_default_paths()

    # ---------------------------------------------------------------
    # BUILD: HEADER
    # ---------------------------------------------------------------
    def _build_header(self):
        header = ctk.CTkFrame(self, corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)

        title_font = ctk.CTkFont(family="Segoe UI", size=22, weight="bold")
        subtitle_font = ctk.CTkFont(family="Segoe UI", size=13)

        ctk.CTkLabel(header, text=APP_TITLE, font=title_font, anchor="w").pack(
            fill="x", padx=20, pady=(16, 0))
        ctk.CTkLabel(header, text=APP_SUBTITLE, font=subtitle_font,
                     text_color=("gray20", "gray70"), anchor="w").pack(
            fill="x", padx=20, pady=(0, 14))

    # ---------------------------------------------------------------
    # BUILD: CONTROL PANEL (checklist path + tombol scan)
    # ---------------------------------------------------------------
    def _build_control_panel(self):
        panel = ctk.CTkFrame(self)
        panel.pack(fill="x", padx=20, pady=(10, 10))

        label_font = ctk.CTkFont(family="Segoe UI", size=14, weight="bold")
        ctk.CTkLabel(panel, text="Pilih Direktori yang Ingin Dipindai",
                     font=label_font, anchor="w").grid(
            row=0, column=0, columnspan=3, sticky="w", padx=15, pady=(12, 6))

        # Scrollable frame berisi checkbox path
        self.paths_frame = ctk.CTkScrollableFrame(panel, height=150)
        self.paths_frame.grid(row=1, column=0, columnspan=3, sticky="ew", padx=15, pady=(0, 10))
        panel.grid_columnconfigure(0, weight=1)

        # Tombol-tombol aksi
        btn_row = ctk.CTkFrame(panel, fg_color="transparent")
        btn_row.grid(row=2, column=0, columnspan=3, sticky="ew", padx=15, pady=(0, 15))

        self.btn_add_path = ctk.CTkButton(
            btn_row, text="+ Tambah Path Manual", width=190,
            fg_color="#3a3a3a", hover_color="#4a4a4a",
            command=self._on_add_manual_path)
        self.btn_add_path.pack(side="left")

        self.btn_scan = ctk.CTkButton(
            btn_row, text="Mulai Scan Direktori", width=220,
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            command=self._on_start_scan)
        self.btn_scan.pack(side="right")

    def _populate_default_paths(self):
        for label, path in DEFAULT_PATHS:
            exists = os.path.isdir(path)
            self._add_path_row(label=label, path=path, checked=exists,
                                enabled=exists, removable=False)

    def _add_path_row(self, label, path, checked=True, enabled=True, removable=True):
        self._path_row_counter += 1
        row_id = self._path_row_counter

        row_frame = ctk.CTkFrame(self.paths_frame, fg_color="transparent")
        row_frame.pack(fill="x", pady=2)

        var = ctk.BooleanVar(value=checked)
        display_text = label if enabled else f"{label}  (tidak ditemukan)"
        text_color = None if enabled else ("gray50", "gray50")

        cb = ctk.CTkCheckBox(row_frame, text=display_text, variable=var)
        if not enabled:
            cb.configure(state="disabled", text_color=text_color)
        cb.pack(side="left", padx=(4, 4), pady=2)

        if removable:
            remove_btn = ctk.CTkButton(
                row_frame, text="✕", width=26, height=24,
                fg_color="#5a2a2a", hover_color="#7a3a3a",
                command=lambda rid=row_id: self._remove_path_row(rid))
            remove_btn.pack(side="right", padx=(4, 8))

        self.path_rows[row_id] = {
            "var": var, "path": path, "label": label,
            "enabled": enabled, "removable": removable, "frame": row_frame,
        }

    def _remove_path_row(self, row_id):
        entry = self.path_rows.pop(row_id, None)
        if entry:
            entry["frame"].destroy()

    def _on_add_manual_path(self):
        folder = filedialog.askdirectory(title="Pilih Folder untuk Dipindai")
        if not folder:
            return
        # Cegah duplikat path yang sudah ada di list
        for entry in self.path_rows.values():
            if os.path.normcase(os.path.normpath(entry["path"])) == os.path.normcase(os.path.normpath(folder)):
                messagebox.showinfo("Info", "Path ini sudah ada di dalam daftar.")
                return
        label = f"Manual: {folder}"
        self._add_path_row(label=label, path=folder, checked=True, enabled=True, removable=True)

    def _get_checked_paths(self):
        """Mengembalikan list (label, path) dari checkbox yang tercentang & aktif."""
        result = []
        for entry in self.path_rows.values():
            if entry["enabled"] and entry["var"].get():
                result.append((entry["label"], entry["path"]))
        return result

    # ---------------------------------------------------------------
    # BUILD: PROGRESS ZONE (progress bar + log console)
    # ---------------------------------------------------------------
    def _build_progress_zone(self):
        zone = ctk.CTkFrame(self)
        zone.pack(fill="x", padx=20, pady=(0, 10))

        top_row = ctk.CTkFrame(zone, fg_color="transparent")
        top_row.pack(fill="x", padx=15, pady=(12, 6))

        self.status_label = ctk.CTkLabel(top_row, text="Siap untuk memindai.", anchor="w")
        self.status_label.pack(side="left")

        self.progress_bar = ctk.CTkProgressBar(zone)
        self.progress_bar.set(0)
        self.progress_bar.pack(fill="x", padx=15, pady=(0, 10))

        self.log_textbox = ctk.CTkTextbox(zone, height=140, font=ctk.CTkFont(family="Consolas", size=12))
        self.log_textbox.pack(fill="x", padx=15, pady=(0, 12))
        self.log_textbox.configure(state="disabled")

    def _append_log(self, text: str):
        self.log_textbox.configure(state="normal")
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log_textbox.insert("end", f"[{timestamp}] {text}\n")
        self.log_textbox.configure(state="disabled")
        self.log_textbox.see("end")

    # ---------------------------------------------------------------
    # BUILD: RESULT ZONE (Treeview preview)
    # ---------------------------------------------------------------
    def _build_result_zone(self):
        zone = ctk.CTkFrame(self)
        zone.pack(fill="both", expand=True, padx=20, pady=(0, 10))

        label_font = ctk.CTkFont(family="Segoe UI", size=14, weight="bold")
        ctk.CTkLabel(zone, text="Hasil Scan (Preview)", font=label_font, anchor="w").pack(
            fill="x", padx=15, pady=(12, 6))

        table_frame = ctk.CTkFrame(zone, fg_color="transparent")
        table_frame.pack(fill="both", expand=True, padx=15, pady=(0, 15))

        style = ttk.Style()
        style.theme_use("clam")
        style.configure("Custom.Treeview",
                         background="#2b2b2b", foreground="white",
                         fieldbackground="#2b2b2b", rowheight=26,
                         borderwidth=0, font=("Segoe UI", 10))
        style.map("Custom.Treeview", background=[("selected", "#1f6aa5")])
        style.configure("Custom.Treeview.Heading",
                         background="#1F3864", foreground="white",
                         font=("Segoe UI", 10, "bold"), relief="flat")
        style.map("Custom.Treeview.Heading", background=[("active", "#274b8a")])

        self.tree = ttk.Treeview(table_frame, columns=PREVIEW_COLUMNS, show="headings",
                                  style="Custom.Treeview")
        col_widths = {
            "No": 45, "Nama Project": 180, "Sumber Server": 200,
            "Framework/Tech": 140, "Ukuran": 90, "Terakhir Dimodifikasi": 140,
            "Jumlah File": 90, "Full Path": 320,
        }
        for col in PREVIEW_COLUMNS:
            self.tree.heading(col, text=col)
            anchor = "center" if col in ("No", "Ukuran", "Jumlah File") else "w"
            self.tree.column(col, width=col_widths.get(col, 120), anchor=anchor)

        vsb = ttk.Scrollbar(table_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(table_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        table_frame.grid_rowconfigure(0, weight=1)
        table_frame.grid_columnconfigure(0, weight=1)

        # Klik ganda -> buka lokasi folder di File Explorer
        self.tree.bind("<Double-1>", self._on_tree_double_click)

    def _on_tree_double_click(self, event):
        item_id = self.tree.focus()
        if not item_id:
            return
        values = self.tree.item(item_id, "values")
        if not values:
            return
        full_path = values[-1]
        if os.path.isdir(full_path):
            try:
                os.startfile(full_path)  # hanya berfungsi di Windows
            except Exception:
                pass

    def _insert_treeview_row(self, row: dict):
        self.tree.insert("", "end", values=(
            row["No"],
            row["Nama Project"],
            row["Sumber Server"],
            row["Framework/Tech"],
            format_size_human(row["Ukuran (MB)"]),
            row["Terakhir Dimodifikasi"],
            row["Jumlah File"],
            row["Full Path"],
        ))

    # ---------------------------------------------------------------
    # BUILD: ACTION ZONE (tombol export)
    # ---------------------------------------------------------------
    def _build_action_zone(self):
        zone = ctk.CTkFrame(self, fg_color="transparent")
        zone.pack(fill="x", padx=20, pady=(0, 16))

        self.btn_export = ctk.CTkButton(
            zone, text="Export Hasil ke Excel (.xlsx)", height=42,
            font=ctk.CTkFont(family="Segoe UI", size=14, weight="bold"),
            fg_color="#1e6b3e", hover_color="#26843f",
            state="disabled", command=self._on_export_excel)
        self.btn_export.pack(fill="x")

    # ---------------------------------------------------------------
    # SCAN: TRIGGER & THREAD WORKER
    # ---------------------------------------------------------------
    def _on_start_scan(self):
        if self.is_scanning:
            return

        checked_paths = self._get_checked_paths()
        if not checked_paths:
            messagebox.showwarning("Peringatan", "Pilih minimal satu direktori untuk dipindai.")
            return

        # Reset state hasil sebelumnya
        self.scan_results = []
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.log_textbox.configure(state="normal")
        self.log_textbox.delete("1.0", "end")
        self.log_textbox.configure(state="disabled")
        self.progress_bar.set(0)

        self.is_scanning = True
        self.btn_scan.configure(state="disabled", text="Sedang Memindai...")
        self.btn_export.configure(state="disabled")
        self.status_label.configure(text="Memindai direktori, mohon tunggu...")

        thread = threading.Thread(target=self._scan_worker, args=(checked_paths,), daemon=True)
        thread.start()
        self.after(80, self._process_queue)

    def _scan_worker(self, checked_paths):
        start_time = time.time()
        results = []
        try:
            all_projects = []  # (source_label, entry_direntry)
            for _label, base_path in checked_paths:
                if not os.path.isdir(base_path):
                    self.log_queue.put(("log", f"[LEWATI] Path tidak ditemukan: {base_path}"))
                    continue
                try:
                    with os.scandir(base_path) as it:
                        subdirs = [e for e in it if e.is_dir(follow_symlinks=False)]
                except PermissionError:
                    self.log_queue.put(("log", f"[ERROR] Akses ditolak: {base_path}"))
                    continue
                except OSError as e:
                    self.log_queue.put(("log", f"[ERROR] {base_path}: {e}"))
                    continue

                source_label = derive_source_label(base_path)
                for e in subdirs:
                    all_projects.append((source_label, e))

            total = len(all_projects)
            self.log_queue.put(("log", f"Ditemukan {total} folder project. Memulai analisis..."))

            for idx, (source_label, entry) in enumerate(all_projects, start=1):
                project_path = entry.path
                project_name = entry.name
                self.log_queue.put(("log", f"[{idx}/{total}] Memindai: {project_name}"))
                try:
                    tech = detect_tech_stack(project_path)
                    size_bytes, file_count, _folder_count, max_mtime = scan_folder_stats(
                        project_path, SKIP_DIR_NAMES)
                    size_mb = round(size_bytes / (1024 * 1024), 2)
                    last_mod = datetime.fromtimestamp(max_mtime).strftime("%Y-%m-%d %H:%M")

                    row = {
                        "No": idx,
                        "Nama Project": project_name,
                        "Sumber Server": source_label,
                        "Full Path": project_path,
                        "Framework/Tech": tech,
                        "Ukuran (MB)": size_mb,
                        "Terakhir Dimodifikasi": last_mod,
                        "Jumlah File": file_count,
                    }
                    results.append(row)
                    self.log_queue.put(("row", row))
                except Exception as e:
                    self.log_queue.put(("log", f"[ERROR] Gagal memindai {project_name}: {e}"))

                if total:
                    self.log_queue.put(("progress", idx / total))

            elapsed = time.time() - start_time
            self.log_queue.put(("log", f"Selesai dalam {elapsed:.1f} detik."))
        except Exception as e:
            self.log_queue.put(("log", f"[FATAL ERROR] {e}"))
        finally:
            self.log_queue.put(("done", results))

    def _process_queue(self):
        try:
            while True:
                kind, payload = self.log_queue.get_nowait()
                if kind == "log":
                    self._append_log(payload)
                elif kind == "progress":
                    self.progress_bar.set(payload)
                elif kind == "row":
                    self._insert_treeview_row(payload)
                elif kind == "done":
                    self._on_scan_done(payload)
        except queue.Empty:
            pass

        if self.is_scanning:
            self.after(80, self._process_queue)

    def _on_scan_done(self, results: list):
        self.scan_results = results
        self.is_scanning = False
        self.btn_scan.configure(state="normal", text="Mulai Scan Direktori")
        self.progress_bar.set(1 if results else 0)

        if results:
            self.status_label.configure(text=f"Selesai! Ditemukan {len(results)} project.")
            self.btn_export.configure(state="normal")
        else:
            self.status_label.configure(text="Selesai. Tidak ada project yang ditemukan.")
            self.btn_export.configure(state="disabled")

    # ---------------------------------------------------------------
    # EXPORT KE EXCEL
    # ---------------------------------------------------------------
    def _on_export_excel(self):
        if not self.scan_results:
            messagebox.showwarning("Peringatan", "Belum ada hasil scan untuk di-export.")
            return

        default_name = f"Laporan_Scan_Directory_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = filedialog.asksaveasfilename(
            title="Simpan Laporan Excel",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel Workbook", "*.xlsx")],
        )
        if not filepath:
            return

        try:
            export_results_to_excel(self.scan_results, filepath)
            messagebox.showinfo("Berhasil", f"Laporan berhasil disimpan ke:\n{filepath}")
            self._append_log(f"Export Excel berhasil: {filepath}")
        except Exception as e:
            messagebox.showerror("Gagal Export", f"Terjadi kesalahan saat export:\n{e}")
            self._append_log(f"[ERROR] Export gagal: {e}")


# =========================================================================
# ENTRY POINT
# =========================================================================

def main():
    app = ScannerApp()
    app.mainloop()


if __name__ == "__main__":
    main()
